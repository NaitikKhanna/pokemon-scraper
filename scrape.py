import time

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


def scrape_page(url):
    try:
        response = requests.get(url, timeout=10)
        response.raise_for_status()
        response.encoding = 'utf-8'
    except requests.RequestException as e:
        print(f"Failed to fetch {url}: {e}")
        return []

    soup = BeautifulSoup(response.text, 'html.parser')
    products = []

    for item in soup.find_all("li", class_="product"):
        title_tag = item.find("h2", class_="woocommerce-loop-product__title")
        price_tag = item.find("span", class_="woocommerce-Price-amount")
        link_tag = item.find("a", href=True)

        if not title_tag or not price_tag or not link_tag:
            print("Skipped one item — missing title, price, or URL")
            continue

        products.append({
            'name': title_tag.text.strip(),
            'price': price_tag.text.strip(),
            'url': link_tag['href']
        })

    return products


def scrape_all_pages(base_url, max_pages=1):
    all_products = []

    for page_num in range(1, max_pages + 1):
        if page_num == 1:
            url = base_url
        else:
            url = f"{base_url}page/{page_num}/"

        print(f"Scraping page {page_num}...")
        products = scrape_page(url)

        if not products:
            break

        all_products.extend(products)
        time.sleep(1)

    return all_products


def clean_data(data):
    df = pd.DataFrame(data)

    if df.empty:
        print("No data scraped.")
        return df

    df['price'] = df['price'].str.replace(r'[^\d.]', '', regex=True)
    df['price'] = df['price'].astype(float)

    df = df.drop_duplicates(subset=['url'])
    df = df.sort_values('price')

    return df


def save_pretty_excel(df, filename):
    df.to_excel(filename, index=False, sheet_name='Products')

    wb = load_workbook(filename)
    ws = wb['Products']

    header_fill = PatternFill(start_color="1C2B3A", end_color="1C2B3A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    for column_cells in ws.columns:
        max_length = max(len(str(cell.value)) for cell in column_cells)
        col_letter = column_cells[0].column_letter
        ws.column_dimensions[col_letter].width = min(max_length + 4, 60)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')

    ws.freeze_panes = "A2"

    wb.save(filename)
    print(f"Saved formatted Excel file: {filename}")


if __name__ == "__main__":
    data = scrape_all_pages(
        "https://scrapeme.live/shop/",
        max_pages=3
    )

    df = clean_data(data)

    if not df.empty:
        save_pretty_excel(df, "scraped_products.xlsx")

    print(df)
    print(f"\nTotal products scraped: {len(df)}")